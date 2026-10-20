"""Bulk employee import from Excel with full validation before anything is written."""

import datetime as dt
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from apps.core.exports import xlsx_bytes
from apps.core.validators import validate_gh_phone, validate_ghana_card, validate_ssnit
from apps.organization.models import Branch, Department, Grade, GradeStep, JobTitle, Location

from .models import BankAccount, Employee, EmployeeCategory
from .services import hire

COLUMNS = [
    "staff_number", "title", "first_name", "middle_name", "last_name", "gender", "date_of_birth", "marital_status",
    "ghana_card_no", "ssnit_number", "tin", "phone", "work_email", "personal_email", "residential_address",
    "digital_address", "hire_date", "probation_end_date", "contract_end_date", "department_code", "branch_code",
    "location", "job_title", "grade_code", "grade_step", "employment_type", "category_code",
    "supervisor_staff_number", "basic_salary", "payment_method", "bank_name", "bank_code", "branch_sort_code",
    "account_number", "account_name", "momo_network", "momo_number",
]
REQUIRED = ["first_name", "last_name", "hire_date"]


def template_bytes():
    example = ["", "mr", "Kwame", "", "Mensah", "male", "15/03/1990", "married", "GHA-123456789-0", "C123456789012",
               "", "0241234567", "k.mensah@example.com", "", "Accra", "GA-123-4567", "01/02/2024", "01/08/2024", "",
               "FIN", "HQ", "Head Office", "Accountant", "G5", "1", "permanent", "senior", "", "4500.00", "bank",
               "GCB Bank", "040", "040101", "1234567890123", "Kwame Mensah", "", ""]
    return xlsx_bytes(COLUMNS, [example], title="Employees")


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"invalid date '{v}' (use DD/MM/YYYY)")


def _str(v):
    return "" if v is None else str(v).strip()


def parse(file):
    wb = load_workbook(file, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], ["The file is empty."]
    header = [_str(h).lower() for h in rows[0]]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        return [], [f"Missing required columns: {', '.join(missing)}"]
    records = []
    for i, r in enumerate(rows[1:], start=2):
        if not any(c not in (None, "") for c in r):
            continue
        records.append((i, {header[j]: r[j] for j in range(min(len(header), len(r))) if header[j]}))
    return records, []


def validate_and_build(records):
    lookups = {
        "department": {d.code.lower(): d for d in Department.objects.all()},
        "branch": {b.code.lower(): b for b in Branch.objects.all()},
        "location": {loc.name.lower(): loc for loc in Location.objects.all()},
        "job_title": {j.name.lower(): j for j in JobTitle.objects.all()},
        "grade": {g.code.lower(): g for g in Grade.objects.all()},
        "category": {c.code.lower(): c for c in EmployeeCategory.objects.all()},
    }
    existing_numbers = set(Employee.objects.values_list("staff_number", flat=True))
    existing_cards = set(Employee.objects.exclude(ghana_card_no="").values_list("ghana_card_no", flat=True))
    seen_numbers, seen_cards = set(), set()
    built, errors = [], []
    for row_no, rec in records:
        errs = []
        data = {}
        try:
            for f in ("title", "first_name", "middle_name", "last_name", "gender", "marital_status", "tin",
                      "work_email", "personal_email", "residential_address", "digital_address", "employment_type"):
                val = _str(rec.get(f))
                if f in ("title", "gender", "marital_status", "employment_type"):
                    val = val.lower()
                if val:
                    data[f] = val
            for f in REQUIRED:
                if not _str(rec.get(f)):
                    errs.append(f"{f} is required")
            for f in ("date_of_birth", "hire_date", "probation_end_date", "contract_end_date"):
                try:
                    d = _date(rec.get(f))
                    if d:
                        data[f] = d
                except ValueError as exc:
                    errs.append(f"{f}: {exc}")
            sn = _str(rec.get("staff_number"))
            if sn:
                if sn in existing_numbers or sn in seen_numbers:
                    errs.append(f"staff number {sn} already exists")
                seen_numbers.add(sn)
                data["staff_number"] = sn
            card = _str(rec.get("ghana_card_no")).upper()
            if card:
                try:
                    validate_ghana_card(card)
                except ValidationError as exc:
                    errs.append(exc.messages[0])
                if card in existing_cards or card in seen_cards:
                    errs.append(f"Ghana Card {card} is already registered")
                seen_cards.add(card)
                data["ghana_card_no"] = card
            ssnit = _str(rec.get("ssnit_number")).upper()
            if ssnit:
                try:
                    validate_ssnit(ssnit)
                except ValidationError as exc:
                    errs.append(exc.messages[0])
                data["ssnit_number"] = ssnit
            phone = _str(rec.get("phone"))
            if phone:
                if phone.isdigit() and len(phone) == 9:
                    phone = "0" + phone  # Excel strips leading zeros
                try:
                    validate_gh_phone(phone)
                except ValidationError as exc:
                    errs.append(exc.messages[0])
                data["phone"] = phone
            for col, key, field in (("department_code", "department", "department"), ("branch_code", "branch", "branch"),
                                    ("location", "location", "location"), ("job_title", "job_title", "job_title"),
                                    ("grade_code", "grade", "grade"), ("category_code", "category", "category")):
                val = _str(rec.get(col)).lower()
                if val:
                    obj = lookups[key].get(val)
                    if obj is None:
                        errs.append(f"{col} '{rec.get(col)}' not found")
                    else:
                        data[field] = obj
            step = _str(rec.get("grade_step"))
            if step and data.get("grade"):
                gs = GradeStep.objects.filter(grade=data["grade"], step=int(float(step))).first()
                if gs is None:
                    errs.append(f"grade step {step} not found for {data['grade'].code}")
                else:
                    data["grade_step"] = gs
            sal = _str(rec.get("basic_salary"))
            if sal:
                try:
                    data["basic_salary"] = Decimal(sal.replace(",", ""))
                except InvalidOperation:
                    errs.append(f"basic_salary '{sal}' is not a number")
            if data.get("date_of_birth") and data.get("hire_date") and data["date_of_birth"] >= data["hire_date"]:
                errs.append("date_of_birth must be before hire_date")
            if data.get("department") and data["department"].cost_center_id:
                data["cost_center"] = data["department"].cost_center
            bank = {k: _str(rec.get(k)) for k in ("payment_method", "bank_name", "bank_code", "branch_sort_code",
                                                   "account_number", "account_name", "momo_network", "momo_number")}
            if bank["account_number"] and "." in bank["account_number"]:
                errs.append("account_number looks like a number in scientific/decimal format - format the column as Text")
        except Exception as exc:  # defensive: never crash the whole import
            errs.append(str(exc))
        if errs:
            errors.append({"row": row_no, "errors": errs})
        else:
            built.append((row_no, data, _str(rec.get("supervisor_staff_number")), bank))
    return built, errors


def run_import(file, user, dry_run=True):
    records, file_errors = parse(file)
    if file_errors:
        return {"ok": False, "created": 0, "errors": [{"row": 0, "errors": file_errors}]}
    built, errors = validate_and_build(records)
    result = {"ok": not errors, "rows": len(records), "valid": len(built), "errors": errors, "created": 0,
              "dry_run": dry_run}
    if errors or dry_run:
        return result
    with transaction.atomic():
        created = {}
        for row_no, data, _sup, bank in built:
            emp = hire(data, user=user)
            created[emp.staff_number] = emp
            method = (bank["payment_method"] or "bank").lower()
            if bank["account_number"] or bank["momo_number"]:
                BankAccount.objects.create(employee=emp, payment_method=method, bank_name=bank["bank_name"],
                                           bank_code=bank["bank_code"], branch_code=bank["branch_sort_code"],
                                           account_number=bank["account_number"],
                                           account_name=bank["account_name"] or emp.full_name,
                                           momo_network=bank["momo_network"].lower(), momo_number=bank["momo_number"])
        for row_no, data, sup, _bank in built:
            if not sup:
                continue
            emp = created.get(data.get("staff_number")) or Employee.objects.filter(
                first_name=data["first_name"], last_name=data["last_name"], hire_date=data["hire_date"]).first()
            boss = created.get(sup) or Employee.objects.filter(staff_number=sup).first()
            if emp and boss and boss != emp:
                emp.supervisor = boss
                emp.save(update_fields=["supervisor"])
                emp.employment_records.filter(is_current=True).update(supervisor=boss)
        result["created"] = len(created)
    return result
