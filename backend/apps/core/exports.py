import csv
import io
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def xlsx_bytes(headers, rows, title="Report", totals_for=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F3D5E")
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    if totals_for and rows:
        total_row = ["TOTAL"] + [""] * (len(headers) - 1)
        for idx in totals_for:
            total_row[idx] = sum(float(r[idx] or 0) for r in rows)
        ws.append(total_row)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r)]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def csv_bytes(headers, rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    for r in rows:
        w.writerow(r)
    return buf.getvalue().encode("utf-8-sig")


def file_response(content: bytes, filename: str, content_type: str):
    resp = HttpResponse(content, content_type=content_type)
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def tabular_response(headers, rows, filename_base, fmt="xlsx", title="Report", totals_for=None):
    if fmt == "csv":
        return file_response(csv_bytes(headers, rows), f"{filename_base}.csv", "text/csv")
    return file_response(xlsx_bytes(headers, rows, title, totals_for), f"{filename_base}.xlsx",
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
